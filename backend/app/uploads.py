"""Chat-only attachments. Not written to CMS.

Office files are extracted to text. Images are stored and sent to the model
as a picture on each turn. The original file stays on disk.
"""

from __future__ import annotations

import base64
import json
import re
import threading
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from pypdf import PdfReader

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "uploads"
INDEX_FILE = DATA_DIR / "index.json"

MAX_BYTES = 15 * 1024 * 1024
MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_CHARS = 16000
MAX_PAGES = 40

_LOCK = threading.Lock()

# kind, stored extension, media type
_BY_EXT = {
    ".pdf": ("pdf", ".pdf", "application/pdf"),
    ".docx": ("word", ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ".xlsx": ("excel", ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ".xls": ("excel", ".xls", "application/vnd.ms-excel"),
    ".doc": ("word", ".doc", "application/msword"),
    ".jpg": ("image", ".jpg", "image/jpeg"),
    ".jpeg": ("image", ".jpg", "image/jpeg"),
    ".png": ("image", ".png", "image/png"),
    ".webp": ("image", ".webp", "image/webp"),
}


class UploadError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _read_index() -> dict[str, Any]:
    if not INDEX_FILE.exists():
        return {"files": []}
    try:
        with INDEX_FILE.open(encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError):
        return {"files": []}
    if not isinstance(data, dict) or not isinstance(data.get("files"), list):
        return {"files": []}
    return data


def _write_index(data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = INDEX_FILE.with_name(INDEX_FILE.name + ".tmp")
    tmp.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    tmp.replace(INDEX_FILE)


def _safe_name(name: str, ext: str) -> str:
    base = Path(name or f"document{ext}").name
    base = re.sub(r"[^\w.\- ()]+", "_", base).strip() or f"document{ext}"
    if Path(base).suffix.lower() not in _BY_EXT:
        base += ext
    return base[:180]


def _clip(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS].rsplit(" ", 1)[0] + "…"
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise UploadError("That file is not a readable PDF.") from exc
    parts: list[str] = []
    for page in reader.pages[:MAX_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = _clip("\n".join(parts))
    if not text:
        raise UploadError(
            "This PDF has no extractable text. It may be a scan — paste the terms you need."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise UploadError("Word support is not installed on this server.") from exc
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise UploadError("That Word file could not be read. Save it as .docx and try again.") from exc
    parts: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _clip("\n".join(parts))
    if not text:
        raise UploadError("That Word file has no extractable text.")
    return text


def _extract_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UploadError("Excel support is not installed on this server.") from exc
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UploadError("That Excel file could not be read. Save it as .xlsx and try again.") from exc
    parts: list[str] = []
    try:
        for sheet in wb.worksheets[:8]:
            parts.append(f"# {sheet.title}")
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 250:
                    parts.append("…")
                    break
                cells = [str(c) for c in row if c is not None and str(c).strip() != ""]
                if cells:
                    parts.append("\t".join(cells))
    finally:
        wb.close()
    text = _clip("\n".join(parts))
    if not text:
        raise UploadError("That spreadsheet has no extractable cells.")
    return text


def _classify(filename: str, data: bytes) -> tuple[str, str, str]:
    ext = Path(filename or "").suffix.lower()
    if ext in _BY_EXT:
        kind, store_ext, media = _BY_EXT[ext]
        if kind == "word" and ext == ".doc":
            raise UploadError("Old .doc files are not supported. Save as .docx and attach that.")
        if kind == "excel" and ext == ".xls":
            raise UploadError("Old .xls files are not supported. Save as .xlsx and attach that.")
        return kind, store_ext, media
    if data.startswith(b"%PDF"):
        return _BY_EXT[".pdf"]
    if data.startswith(b"\xff\xd8\xff"):
        return _BY_EXT[".jpg"]
    if data.startswith(b"\x89PNG"):
        return _BY_EXT[".png"]
    raise UploadError(
        "Attach a PDF, Word (.docx), Excel (.xlsx), or image (.jpg / .png)."
    )


def save_pdf(filename: str, data: bytes) -> dict[str, Any]:
    """Save any supported chat attachment. Name kept for the existing route."""
    if not data:
        raise UploadError("The file was empty.")
    if len(data) > MAX_BYTES:
        raise UploadError("File is too large (15 MB limit).")

    kind, store_ext, media = _classify(filename, data)
    if kind == "image" and len(data) > MAX_IMAGE_BYTES:
        raise UploadError("Images must be 5 MB or smaller.")

    if kind == "pdf":
        if not data.startswith(b"%PDF"):
            raise UploadError("That file is not a valid PDF.")
        text = _extract_pdf(data)
    elif kind == "word":
        text = _extract_docx(data)
    elif kind == "excel":
        text = _extract_xlsx(data)
    else:
        text = "(image attached — look at the picture sent with the user's message)"

    file_id = f"att-{uuid.uuid4().hex[:12]}"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path = DATA_DIR / f"{file_id}{store_ext}"
    path.write_bytes(data)

    record = {
        "id": file_id,
        "filename": _safe_name(filename, store_ext),
        "kind": kind,
        "ext": store_ext,
        "media_type": media,
        "bytes": len(data),
        "created_at": _now(),
    }
    with _LOCK:
        index = _read_index()
        index["files"].append(record)
        _write_index(index)
        (DATA_DIR / f"{file_id}.txt").write_text(text, encoding="utf-8")
    return {
        "id": record["id"],
        "filename": record["filename"],
        "kind": kind,
        "chars": len(text),
    }


def get(file_id: str) -> dict[str, Any] | None:
    if not file_id or not str(file_id).startswith("att-"):
        return None
    text_path = DATA_DIR / f"{file_id}.txt"
    if not text_path.is_file():
        return None
    meta = {
        "filename": f"{file_id}.pdf",
        "kind": "pdf",
        "ext": ".pdf",
        "media_type": "application/pdf",
    }
    with _LOCK:
        for row in _read_index().get("files") or []:
            if row.get("id") == file_id:
                meta.update({
                    "filename": row.get("filename") or meta["filename"],
                    "kind": row.get("kind") or meta["kind"],
                    "ext": row.get("ext") or meta["ext"],
                    "media_type": row.get("media_type") or meta["media_type"],
                })
                break
    path = DATA_DIR / f"{file_id}{meta['ext']}"
    if not path.is_file():
        for cand in DATA_DIR.glob(f"{file_id}.*"):
            if cand.suffix.lower() != ".txt":
                path = cand
                break
        else:
            path = None
    return {
        "id": file_id,
        "filename": meta["filename"],
        "kind": meta["kind"],
        "media_type": meta["media_type"],
        "text": text_path.read_text(encoding="utf-8"),
        "path": path if path and path.is_file() else None,
    }


def public_meta(file_id: str | None) -> dict[str, str] | None:
    rec = get(file_id) if file_id else None
    if not rec:
        return None
    return {"id": rec["id"], "filename": rec["filename"], "kind": rec["kind"]}


def with_image_on_last_user(history: list[dict[str, Any]], attachment: dict[str, Any]) -> list[dict[str, Any]]:
    """Put the attached image on the latest user turn so Claude can see it."""
    if not history or attachment.get("kind") != "image" or not attachment.get("path"):
        return history
    last = history[-1]
    if last.get("role") != "user" or not isinstance(last.get("content"), str):
        return history
    try:
        raw = Path(attachment["path"]).read_bytes()
        b64 = base64.standard_b64encode(raw).decode("ascii")
    except OSError:
        return history
    media = attachment.get("media_type") or "image/jpeg"
    rebuilt = list(history[:-1])
    rebuilt.append({
        "role": "user",
        "content": [
            {
                "type": "image",
                "source": {"type": "base64", "media_type": media, "data": b64},
            },
            {"type": "text", "text": last["content"]},
        ],
    })
    return rebuilt
